"""Leitura das planilhas de referência (Filiais, Fornecedores, Pedidos de Compra)."""
from pathlib import Path
from typing import Optional
import openpyxl

from app.config import settings
from app.core.exceptions import ReferenciaNaoEncontrada


# Cache simples em memória (invalidado quando o arquivo muda)
_cache: dict = {}


def _normalizar_cnpj(valor) -> str:
    """Mantém apenas dígitos do CNPJ."""
    if valor is None:
        return ""
    return "".join(c for c in str(valor) if c.isdigit())


def _strip(valor) -> str:
    return str(valor).strip() if valor is not None else ""


def _carregar_xlsx(path: Path) -> list[dict]:
    """Lê uma planilha .xlsx e retorna lista de dicts (cabeçalho da primeira linha)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    out: list[dict] = []
    for row in rows:
        if all(c is None for c in row):
            continue
        out.append(dict(zip(headers, row)))
    wb.close()
    return out


def _get_cached(tipo: str, path: Path):
    """Retorna dados cacheados ou recarrega se o arquivo mudou."""
    if not path.exists():
        raise ReferenciaNaoEncontrada(f"Planilha {tipo} não encontrada em {path}")
    mtime = path.stat().st_mtime
    cached = _cache.get(tipo)
    if cached and cached["mtime"] == mtime:
        return cached["data"]
    data = _carregar_xlsx(path)
    _cache[tipo] = {"mtime": mtime, "data": data}
    return data


# ---------------------------------------------------------------------------
# Validação de planilha (no upload)
# ---------------------------------------------------------------------------

COLUNAS_OBRIGATORIAS = {
    "filiais": {"NUMERO DA FILIAL", "CNPJ"},
    "fornecedores": {"NUMERO DO FORNECEDOR", "NOME", "CNPJ"},
    "pedidos_compra": {
        "PC NUM", "FILIAL_COMPRADORA", "FORNECEDOR_COD", "FORNECEDORCNPJ",
        "ITEM_DES_PRD", "ITEM_PRC_UNT", "OBS_PEDIDO",
    },
}


def validar_planilha(path: Path, tipo: str) -> dict:
    """Valida que a planilha contém as colunas esperadas e retorna info."""
    if tipo not in COLUNAS_OBRIGATORIAS:
        raise ValueError(f"Tipo desconhecido: {tipo}")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    headers = [str(c).strip() if c is not None else "" for c in next(ws.iter_rows(values_only=True))]
    wb.close()
    obrigatorias = COLUNAS_OBRIGATORIAS[tipo]
    faltando = obrigatorias - set(headers)
    if faltando:
        raise ValueError(f"Colunas obrigatórias faltando: {faltando}")
    # contagem rápida (reabrindo só pra contar)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    total = ws.max_row - 1 if ws.max_row else 0
    wb.close()
    return {"linhas": total, "colunas": headers}


# ---------------------------------------------------------------------------
# Listagens / consultas
# ---------------------------------------------------------------------------

def listar_filiais() -> list[dict]:
    path = settings.referencias_dir / settings.arquivo_filiais
    rows = _get_cached("filiais", path)
    return [{"codigo": _strip(r["NUMERO DA FILIAL"]), "cnpj": _normalizar_cnpj(r["CNPJ"])} for r in rows]


def listar_fornecedores(busca: Optional[str] = None, limite: int = 50) -> list[dict]:
    path = settings.referencias_dir / settings.arquivo_fornecedores
    rows = _get_cached("fornecedores", path)
    out: list[dict] = []
    busca_norm = (busca or "").upper().strip()
    for r in rows:
        nome = _strip(r["NOME"]).upper()
        cnpj = _normalizar_cnpj(r["CNPJ"])
        if busca_norm and busca_norm not in nome and busca_norm not in cnpj:
            continue
        out.append({
            "codigo": _strip(r["NUMERO DO FORNECEDOR"]),
            "nome": _strip(r["NOME"]),
            "cnpj": cnpj,
        })
        if len(out) >= limite:
            break
    return out


def listar_pedidos_compra(limite: int = 50) -> list[dict]:
    path = settings.referencias_dir / settings.arquivo_pedidos_compra
    rows = _get_cached("pedidos_compra", path)
    return [_pc_to_dict(r) for r in rows[:limite]]


def _pc_to_dict(r: dict) -> dict:
    return {
        "pc_num": _strip(r.get("PC NUM")),
        "status": _strip(r.get("STATUS_PEDIDO")),
        "filial": _strip(r.get("FILIAL_COMPRADORA")),
        "fornecedor_cod": _strip(r.get("FORNECEDOR_COD")),
        "fornecedor_cnpj": _normalizar_cnpj(r.get("FORNECEDORCNPJ")),
        "fornecedor_nome": _strip(r.get("FORNECEDOR_NOME")),
        "item_descricao": _strip(r.get("ITEM_DES_PRD")),
        "item_prc_unt": float(r.get("ITEM_PRC_UNT") or 0),
        "obs_pedido": _strip(r.get("OBS_PEDIDO")) or None,
        "dt_emissao": str(r.get("DT EMISSAO PC") or ""),
    }


# ---------------------------------------------------------------------------
# Lookups específicos (usados pelos parsers / matcher)
# ---------------------------------------------------------------------------

# Regras fixas para CNPJs de matriz com múltiplas filiais
FILIAL_FIXA_POR_CNPJ_MATRIZ = {
    "73305997000161": "01001",
    "67844183000100": "03001",
}


def filial_por_cnpj(cnpj: str) -> Optional[str]:
    """Retorna o código da filial para o CNPJ fornecido. Aplica regras fixas de matriz."""
    cnpj_norm = _normalizar_cnpj(cnpj)
    if cnpj_norm in FILIAL_FIXA_POR_CNPJ_MATRIZ:
        return FILIAL_FIXA_POR_CNPJ_MATRIZ[cnpj_norm]
    for f in listar_filiais():
        if f["cnpj"] == cnpj_norm:
            return f["codigo"]
    return None


def fornecedor_por_cnpj(cnpj: str) -> Optional[dict]:
    """Retorna {codigo, nome, cnpj} ou None."""
    cnpj_norm = _normalizar_cnpj(cnpj)
    if not cnpj_norm:
        return None
    path = settings.referencias_dir / settings.arquivo_fornecedores
    rows = _get_cached("fornecedores", path)
    for r in rows:
        if _normalizar_cnpj(r["CNPJ"]) == cnpj_norm:
            return {
                "codigo": _strip(r["NUMERO DO FORNECEDOR"]),
                "nome": _strip(r["NOME"]),
                "cnpj": cnpj_norm,
            }
    return None


def pedidos_compra_por_fornecedor_filial(fornecedor_cnpj: str, filial: str) -> list[dict]:
    """Retorna todos os PCs (linhas de itens) que casam com fornecedor+filial."""
    cnpj_norm = _normalizar_cnpj(fornecedor_cnpj)
    path = settings.referencias_dir / settings.arquivo_pedidos_compra
    rows = _get_cached("pedidos_compra", path)
    out: list[dict] = []
    for r in rows:
        if _normalizar_cnpj(r.get("FORNECEDORCNPJ")) != cnpj_norm:
            continue
        if _strip(r.get("FILIAL_COMPRADORA")) != filial:
            continue
        out.append(_pc_to_dict(r))
    return out
