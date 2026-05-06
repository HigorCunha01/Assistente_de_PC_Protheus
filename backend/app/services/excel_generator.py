"""Gera Excel a partir das fichas consolidadas.

Layout: 2 abas
- "Pedidos": 1 linha por ficha (cabeçalho + sugestão de PC)
- "Itens": 1 linha por item, com FK pra ficha (FORNECEDOR_NUMS)
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.ficha import FichaPedido


HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def gerar_excel(fichas: list[FichaPedido], destino: Path) -> Path:
    wb = openpyxl.Workbook()

    # Aba 1: Pedidos
    ws_p = wb.active
    ws_p.title = "Pedidos"
    cabec_pedidos = [
        "Nome do arquivo",
        "Filial",
        "CNPJ Filial",
        "Fornecedor (cód)",
        "Fornecedor (nome)",
        "CNPJ Fornecedor",
        "Tipo",
        "Documentos (tipo)",
        "Números dos docs",
        "PC para cópia",
        "OBS Original do PC",
        "OBS Modificada (próximo mês)",
        "Status extração",
        "Observações",
    ]
    _escrever_cabec(ws_p, cabec_pedidos)

    for i, f in enumerate(fichas, start=2):
        ws_p.cell(row=i, column=1, value=", ".join(f.nomes_arquivos_padrao))
        ws_p.cell(row=i, column=2, value=f.filial_codigo or "(não encontrado)")
        ws_p.cell(row=i, column=3, value=f.filial_cnpj or "(não encontrado)")
        ws_p.cell(row=i, column=4, value=f.fornecedor_codigo or "(não encontrado)")
        ws_p.cell(row=i, column=5, value=f.fornecedor_nome or "(não encontrado)")
        ws_p.cell(row=i, column=6, value=f.fornecedor_cnpj or "(não encontrado)")
        ws_p.cell(row=i, column=7, value=f.tipo_pedido or "(não encontrado)")
        ws_p.cell(row=i, column=8, value=f.documentos_tipo or "")
        ws_p.cell(row=i, column=9, value=", ".join(f.numeros_documentos))
        ws_p.cell(row=i, column=10, value=f.sugestao_pc_num or "")
        ws_p.cell(row=i, column=11, value=f.sugestao_pc_obs_original or "")
        ws_p.cell(row=i, column=12, value=f.sugestao_pc_obs_modificada or "")
        ws_p.cell(row=i, column=13, value=f.status_extracao)
        ws_p.cell(row=i, column=14, value=" | ".join(f.observacoes))

    _ajustar_largura(ws_p, cabec_pedidos)

    # Aba 2: Itens
    ws_i = wb.create_sheet("Itens")
    cabec_itens = [
        "Filial",
        "Fornecedor (cód)",
        "Fornecedor (nome)",
        "Documentos",
        "Item nº",
        "Descrição",
        "Quantidade",
        "Valor unitário",
        "Valor total",
        "Unitário calculado?",
    ]
    _escrever_cabec(ws_i, cabec_itens)

    linha = 2
    for f in fichas:
        for n, item in enumerate(f.itens, start=1):
            ws_i.cell(row=linha, column=1, value=f.filial_codigo or "")
            ws_i.cell(row=linha, column=2, value=f.fornecedor_codigo or "")
            ws_i.cell(row=linha, column=3, value=f.fornecedor_nome or "")
            ws_i.cell(row=linha, column=4, value=", ".join(f.numeros_documentos))
            ws_i.cell(row=linha, column=5, value=n)
            ws_i.cell(row=linha, column=6, value=item.descricao)
            ws_i.cell(row=linha, column=7, value=float(item.quantidade))
            ws_i.cell(row=linha, column=8, value=float(item.valor_unitario))
            ws_i.cell(row=linha, column=9, value=float(item.quantidade) * float(item.valor_unitario))
            ws_i.cell(row=linha, column=10, value="sim" if item.valor_unitario_calculado else "não")
            linha += 1

    _ajustar_largura(ws_i, cabec_itens)

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(destino))
    return destino


def _escrever_cabec(ws, cabec: list[str]) -> None:
    for col, header in enumerate(cabec, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _ajustar_largura(ws, cabec: list[str]) -> None:
    for col_idx, header in enumerate(cabec, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(header)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(60, max(12, max_len + 2))
