r"""Gera uma planilha simples com Filial, Cod Fornecedor e Pedido p/ copia.

Uso:
    cd C:\Contas\Contas 2026\05 - MAIO
    python C:\dev\protheus-pc\scripts\gerar_pcs_copia.py

Se nenhuma pasta for informada, o script lê todos os PDFs da pasta atual.
"""
from __future__ import annotations

import argparse
import os
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


SCRIPT_DIR = Path(__file__).resolve().parent
REPO_DIR = SCRIPT_DIR.parent
BACKEND_DIR = REPO_DIR / "backend"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Le todos os PDFs de uma pasta e gera Excel com Filial, Cod Fornecedor e Pedido p/ copia."
    )
    parser.add_argument(
        "pasta",
        nargs="?",
        default=".",
        help="Pasta onde estao os PDFs. Se omitido, usa a pasta atual do terminal.",
    )
    parser.add_argument(
        "-o",
        "--saida",
        default="pc_para_copia.xlsx",
        help="Nome/caminho do Excel de saida. Padrao: pc_para_copia.xlsx",
    )
    parser.add_argument(
        "-r",
        "--recursivo",
        action="store_true",
        help="Procura PDFs tambem nas subpastas.",
    )
    args = parser.parse_args()

    pasta = Path(args.pasta).resolve()
    if not pasta.exists() or not pasta.is_dir():
        print(f"Pasta nao encontrada: {pasta}")
        return 1

    saida = Path(args.saida)
    if not saida.is_absolute():
        saida = pasta / saida

    _configurar_backend()

    try:
        from app.schemas.ficha import FichaPedido
        from app.services.pc_matcher import sugerir_pedido_compra
        from app.services.pdf.extractor import extrair_documento
    except ModuleNotFoundError as exc:
        print(f"Dependencia Python nao encontrada: {exc.name}")
        print("Instale as dependencias com:")
        print(r"  pip install -r C:\dev\protheus-pc\backend\requirements.txt")
        return 1

    _configurar_tesseract()

    padrao = pasta.rglob("*") if args.recursivo else pasta.iterdir()
    arquivos = sorted({p for p in padrao if p.is_file() and p.suffix.lower() == ".pdf"})

    if not arquivos:
        print(f"Nenhum PDF encontrado em: {pasta}")
        return 1

    rows: list[tuple[str, str, str, str, float | None]] = []
    erros: list[tuple[str, str]] = []

    for pdf in arquivos:
        try:
            dados = extrair_documento(pdf, nome_original=pdf.name)
            ficha = FichaPedido(**dados)
            sugerir_pedido_compra(ficha)
            valor_total = _valor_total_documento(ficha)
            rows.append(
                (
                    pdf.name,
                    ficha.filial_codigo or "",
                    ficha.fornecedor_codigo or "",
                    ficha.sugestao_pc_num or "",
                    valor_total,
                )
            )
            print(f"OK   {pdf.name}")
        except Exception as exc:
            erros.append((pdf.name, str(exc)))
            rows.append((pdf.name, "", "", "", None))
            print(f"ERRO {pdf.name}: {exc}")

    _gerar_excel(rows, saida)

    print()
    print(f"Excel gerado: {saida}")
    print(f"PDFs processados: {len(arquivos)}")
    if erros:
        print(f"Arquivos com erro: {len(erros)}")
    return 0


def _valor_total_documento(ficha) -> float:
    return sum(float(item.quantidade) * float(item.valor_unitario) for item in ficha.itens)


def _gerar_excel(rows: list[tuple[str, str, str, str, float | None]], saida: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PC para copia"

    headers = ["Nome do arquivo", "Filial", "Cód Fornecedor", "Pedido p/ cópia", "Valor total do documento"]
    ws.append(headers)

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row in rows:
        ws.append(row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 24

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=5).number_format = '#,##0.00'

    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)


def _configurar_backend() -> None:
    if BACKEND_DIR.exists():
        sys.path.insert(0, str(BACKEND_DIR))
        os.environ.setdefault("REFERENCIAS_DIR", str(REPO_DIR / "data" / "referencias"))
        os.environ.setdefault("TMP_DIR", str(REPO_DIR / "data" / "tmp"))
        return

    # Caminho usado dentro do container backend do Docker Compose.
    container_backend = Path("/app")
    if (container_backend / "app").exists():
        sys.path.insert(0, str(container_backend))


def _configurar_tesseract() -> None:
    if shutil.which("tesseract"):
        return

    caminhos_comuns = [
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    ]
    for caminho in caminhos_comuns:
        if caminho.exists():
            try:
                import pytesseract

                pytesseract.pytesseract.tesseract_cmd = str(caminho)
            except ModuleNotFoundError:
                pass
            return

    print("AVISO: Tesseract OCR nao encontrado.")
    print("PDFs escaneados/imagem podem sair sem filial, fornecedor e valor.")
    print("Instale com: winget install UB-Mannheim.TesseractOCR")
    print()


if __name__ == "__main__":
    raise SystemExit(main())
