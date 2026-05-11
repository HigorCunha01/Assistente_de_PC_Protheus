r"""Gera Excel com PC para copia lendo todos os PDFs de uma pasta.

Uso comum:
    cd C:\Contas\Contas 2026\05 - MAIO
    python C:\dev\protheus-pc\pc-copia-script\gerar_pcs_copia.py

O script usa as planilhas em:
    <pasta deste projeto>\referencias
"""
from __future__ import annotations

import argparse
import os
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


PROJECT_DIR = Path(__file__).resolve().parent
REFERENCIAS_DIR = PROJECT_DIR / "referencias"
TMP_DIR = PROJECT_DIR / "data" / "tmp"

sys.path.insert(0, str(PROJECT_DIR))
os.environ["REFERENCIAS_DIR"] = str(REFERENCIAS_DIR)
os.environ["TMP_DIR"] = str(TMP_DIR)

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Le todos os PDFs de uma pasta e gera Excel com filial, fornecedor, PC para copia e valor."
    )
    parser.add_argument(
        "pasta",
        nargs="?",
        default=".",
        help="Pasta onde estao os PDFs. Se omitido, usa a pasta atual do terminal.",
    )
    parser.add_argument(
        "-o",
        "--saida",
        default="pc_para_copia.xlsx",
        help="Nome/caminho do Excel de saida. Padrao: pc_para_copia.xlsx",
    )
    parser.add_argument(
        "-r",
        "--recursivo",
        action="store_true",
        help="Procura PDFs tambem nas subpastas.",
    )
    args = parser.parse_args()

    pasta = Path(args.pasta).resolve()
    if not pasta.exists() or not pasta.is_dir():
        print(f"Pasta nao encontrada: {pasta}")
        return 1

    saida = Path(args.saida)
    if not saida.is_absolute():
        saida = pasta / saida

    _configurar_tesseract()
    _validar_referencias()

    try:
        from app.schemas.ficha import FichaPedido
        from app.services.pc_matcher import sugerir_pedido_compra
        from app.services.pdf.extractor import extrair_documento
    except ModuleNotFoundError as exc:
        print(f"Dependencia Python nao encontrada: {exc.name}")
        print("Instale as dependencias com:")
        print(f'  python -m pip install -r "{PROJECT_DIR / "requirements.txt"}"')
        return 1

    TMP_DIR.mkdir(parents=True, exist_ok=True)
    arquivos = _listar_pdfs(pasta, recursivo=args.recursivo)
    if not arquivos:
        print(f"Nenhum PDF encontrado em: {pasta}")
        return 1

    rows: list[tuple[str, str, str, str, float | None]] = []
    erros: list[tuple[str, str]] = []

    for pdf in arquivos:
        try:
            dados = extrair_documento(pdf, nome_original=pdf.name)
            ficha = FichaPedido(**dados)
            sugerir_pedido_compra(ficha)
            rows.append(
                (
                    pdf.name,
                    ficha.filial_codigo or "",
                    ficha.fornecedor_codigo or "",
                    ficha.sugestao_pc_num or "",
                    _valor_total_documento(ficha),
                )
            )
            print(f"OK   {pdf.name}")
        except Exception as exc:
            erros.append((pdf.name, str(exc)))
            rows.append((pdf.name, "", "", "", None))
            print(f"ERRO {pdf.name}: {exc}")

    _gerar_excel(rows, saida)

    print()
    print(f"Excel gerado: {saida}")
    print(f"PDFs processados: {len(arquivos)}")
    if erros:
        print(f"Arquivos com erro: {len(erros)}")
    return 0


def _listar_pdfs(pasta: Path, recursivo: bool) -> list[Path]:
    padrao = pasta.rglob("*") if recursivo else pasta.iterdir()
    return sorted({p for p in padrao if p.is_file() and p.suffix.lower() == ".pdf"})


def _validar_referencias() -> None:
    esperados = ["Filiais.xlsx", "Fornecedores.xlsx", "Pedidos_de_Compra.xlsx"]
    faltando = [nome for nome in esperados if not (REFERENCIAS_DIR / nome).exists()]
    if faltando:
        print("AVISO: faltam planilhas de referencia:")
        for nome in faltando:
            print(f"  - {REFERENCIAS_DIR / nome}")
        print()


def _valor_total_documento(ficha) -> float:
    return sum(float(item.quantidade) * float(item.valor_unitario) for item in ficha.itens)


def _gerar_excel(rows: list[tuple[str, str, str, str, float | None]], saida: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PC para copia"

    headers = ["Nome do arquivo", "Filial", "Cod Fornecedor", "Pedido p/ copia", "Valor total do documento"]
    ws.append(headers)

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row in rows:
        ws.append(row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 24

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=5).number_format = '#,##0.00'

    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)


def _configurar_tesseract() -> None:
    if shutil.which("tesseract"):
        return

    caminhos_comuns = [
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    ]
    for caminho in caminhos_comuns:
        if caminho.exists():
            try:
                import pytesseract

                pytesseract.pytesseract.tesseract_cmd = str(caminho)
            except ModuleNotFoundError:
                pass
            return

    print("AVISO: Tesseract OCR nao encontrado.")
    print("PDFs escaneados/imagem podem sair sem filial, fornecedor e valor.")
    print("Instale com: winget install UB-Mannheim.TesseractOCR")
    print()


if __name__ == "__main__":
    raise SystemExit(main())
